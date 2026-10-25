#!/usr/bin/env python3
"""Regression checks for the generated ZIN STONE DOCX/PDF layout.
Run after regenerating the preview via the updated company adapter or demo.
Uses python-docx and pymupdf, not production .NET/PDFShift.
"""
from pathlib import Path
from docx import Document
import fitz

ROOT=Path(__file__).resolve().parents[1]
TEMPLATE=ROOT/'backend/Templates/ZIN_STONE_Quotation.docx'
DEMO=ROOT/'preview/ZIN_STONE_Quote_Demo.docx'
PDF=ROOT/'preview/ZIN_STONE_Quote_Demo.pdf'
EXPECTED=['النوع','المقاس','الحمولة','سعر الألف']
WIDTHS=[2760,2260,1570,1932]

for path in (TEMPLATE,DEMO):
    doc=Document(path)
    assert len(doc.tables)==1, f'product table count: {path}'
    table=doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == EXPECTED, f'bad headers in {path}'
    for index,c in enumerate(table.rows[0].cells):
        assert c.width.twips==WIDTHS[index],f'bad width col {index} in {path}'
    rows=table.rows[1:]
    for r in rows:
        for c in r.cells:
            p=c.paragraphs[0]
            assert int(p.alignment)==1, 'cell not centered'
            assert all(run.bold for run in p.runs if run.text.strip()), 'cell text not bold'
    title=[p for p in doc.paragraphs if 'المدير المسؤل' in p.text]
    name=[p for p in doc.paragraphs if 'يوسف زين الدين' in p.text]
    assert len(title)==len(name)==1, 'signature missing or split'
    assert len(title[0].runs)==len(name[0].runs)==1, 'signature has old scattered runs'
    assert name[0].text=='يوسف زين الدين',repr(name[0].text)
    assert title[0].text=='المدير المسؤل',repr(title[0].text)

assert len(Document(DEMO).tables[0].rows)==3, 'demo must have 2 product lines'
with fitz.open(PDF) as pdf:
    assert len(pdf)==1, f'demo must be one page, got {len(pdf)}'
    page=pdf[0]
    spans=[(sp['text'].strip(),sp['bbox']) for b in page.get_text('dict')['blocks']
           for line in b.get('lines',[]) for sp in line.get('spans',[])]
    names=[(t,b) for t,b in spans if t=='يوسف زين الدين']
    titles=[(t,b) for t,b in spans if t=='المدير المسؤل']
    assert len(names)==len(titles)==1, 'PDF signature broken into parts'
    assert names[0][1][0]<page.rect.width*0.24, 'name must be physically at left of page'
    assert titles[0][1][0]<page.rect.width*0.24, 'title must be physically at left'
    assert abs(names[0][1][0]-titles[0][1][0])<4, 'signature rows not left-aligned'
    assert names[0][1][1]>titles[0][1][1], 'name must be below title'
    assert all(any(k in t for t,_ in spans) for k in ('4200','3500','بلوك خرساني','طوب مصمت')), 'missing product data'
print('PASS: four fixed aligned columns, bold centered cells, two intact left-side signature lines, one-page PDF')
